#!/usr/bin/env python3
"""Read the master workbook's own layout, so nothing about it is typed.

The rate card export writes values into a COPY of the master rather than
building a sheet, which is the only way the styling, merges, column
widths, row heights, logo drawing and print area survive. To do that it
has to know which row each of the 45 rates prints on.

That is a property of the customer's file, not of this repository, so it
is read out of the file. A new master with a row inserted is picked up by
running this again and committing the result, rather than by somebody
noticing that one rate is a row out.

    python3 scripts/rate-card-sheet-map.py
"""
import json
import sys
from openpyxl import load_workbook

MASTER = 'docs/source/rate_cards/master/KNDS UK - Customer Rates 2026.xlsx'
MODEL = 'docs/source/rate_cards/rate-model.json'
OUT = 'docs/source/rate_cards/sheet-map.json'
STATIC = 'docs/source/rate_cards/sheet-static.json'


def main() -> int:
    ws = load_workbook(MASTER)['Costing Info']
    model = json.load(open(MODEL))

    # ---- Section AND item, never item alone ----
    #
    # Seven rates are called "In hours" or "Out of Hours": trailers, HGVs
    # and the bodyshop each have one, and so does callout. Matching on the
    # label alone puts all of them on the first row that carries it, which
    # writes the callout rate into the trailer row and leaves five rates
    # unwritten. It looked right until the export was compared against the
    # master cell by cell.
    #
    # Column A carries the section, set on the first row of each one and
    # blank afterwards, so it is carried down as the sheet is read.
    by_key = {}
    section = ''
    for r in range(10, 70):
        head = ws.cell(r, 1).value
        if head and str(head).strip():
            section = str(head).strip()
        label = ws.cell(r, 2).value
        if label and str(label).strip():
            by_key.setdefault((section, str(label).strip()), r)

    def row_for(sec: str, item: str):
        if (sec, item) in by_key:
            return by_key[(sec, item)]
        # The master qualifies four labels the kit shortens, for example
        # 'HGV Lane Fee (Max per DVSA £70)'. A prefix is the match, and it
        # is still looked for inside the right section.
        for (s2, name), row in by_key.items():
            if s2 == sec and name.startswith(item):
                return row
        # A section the kit names differently from the workbook. Fall back
        # to a unique label anywhere, and only when it IS unique, so a
        # collision is reported rather than guessed at.
        hits = [row for (_, name), row in by_key.items()
                if name == item or name.startswith(item)]
        return hits[0] if len(set(hits)) == 1 else None

    rows, missing = {}, []
    for i, rate in enumerate(model['rates']):
        rid = f'r{i + 1:02d}'
        row = row_for(rate['section'].strip(), rate['item'].strip())
        if row is None:
            missing.append((rid, f"{rate['section']} / {rate['item']}"))
        else:
            rows[rid] = row

    if missing:
        for rid, item in missing:
            print(f'  {rid}: no row in the master for "{item}"', file=sys.stderr)
        print('\n  The master and the model disagree. Fix that before exporting anything.',
              file=sys.stderr)
        return 1

    # Row 10 names the columns, so which letter is which is read too.
    header = {str(ws.cell(10, c).value or '').strip(): ws.cell(10, c).column_letter
              for c in range(3, 10)}
    columns = {
        'single': header.get('Price', 'C'),
        'axle': [header.get(f'{n}-axle', chr(ord('D') + n - 1)) for n in (1, 2, 3, 4)],
    }

    json.dump({'rows': rows, 'columns': columns}, open(OUT, 'w'), indent=2)
    print(f'  {len(rows)} rates mapped to rows, columns {columns["single"]} '
          f'and {"".join(columns["axle"])}, written to {OUT}')

    # ---- The words the master prints that no card supplies ----
    #
    # The export writes values into a copy of the master, so every label
    # on the sheet comes along for free: 'Main Contact', 'Hourly Rate -
    # Trailers', 'Price', '1-axle', 'FleetSmart+ Inclusions', 'All work
    # MUST be authorised prior to commencement'.
    #
    # A PDF drawn from the values alone has none of them. It is a grid of
    # prices with nothing saying what any of them is for, which is not
    # the same document the customer gets as a spreadsheet, and the
    # scope's whole point is that the two are one document.
    #
    # So they are read out of the master here, with the master's own bold
    # and alignment, and every renderer draws them.
    #
    # Which cells are labels and which are one customer's data is a
    # property of the sheet's shape: column A and column B name things,
    # row 10 names the columns, K8 is the note over the inclusions
    # panel. The exceptions are A2, the customer's name, and B3 to B8,
    # their contact block. `check:rate-card-pdf` asserts this set never
    # overlaps what a card writes, so a new master that moves somebody's
    # data into column B fails rather than printing KNDS's details on
    # another customer's card.
    labels = []
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        for cell in row:
            text = cell.value
            if text is None or str(text).strip() == '':
                continue
            col, r = cell.column_letter, cell.row
            is_label = (
                (col == 'A' and r >= 3)
                or (col == 'B' and not 3 <= r <= 8)
                or r == 10
                or cell.coordinate == 'K8'
            )
            if not is_label:
                continue
            labels.append({
                'at': cell.coordinate,
                'text': str(text).rstrip(),
                'bold': bool(cell.font and cell.font.bold),
                'align': cell.alignment.horizontal if cell.alignment else None,
            })

    json.dump(labels, open(STATIC, 'w'), indent=2, ensure_ascii=False)
    print(f'  {len(labels)} of the master\'s own labels read out, written to {STATIC}')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
